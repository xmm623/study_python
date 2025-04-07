from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import datetime

# 创建一个工作簿
wb = Workbook()
ws = wb.active
ws.title = "阅读APP首页短文展示策略优化测试"

# 设置样式
header_font = Font(name='Arial', bold=True, size=12)
header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
border = Border(
    left=Side(border_style="thin", color="000000"),
    right=Side(border_style="thin", color="000000"),
    top=Side(border_style="thin", color="000000"),
    bottom=Side(border_style="thin", color="000000")
)

# 设置列宽
column_widths = {
    'A': 8,   # 序号
    'B': 25,  # 测试项
    'C': 30,  # 测试步骤
    'D': 30,  # 期望结果
    'E': 15,  # 测试结果
    'F': 15   # 备注
}

for col, width in column_widths.items():
    ws.column_dimensions[col].width = width

# 设置标题行
headers = ["序号", "测试项", "测试步骤", "期望结果", "测试结果", "备注"]
for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.value = header
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = border

# 定义测试用例数据
test_cases = [
    {
        "序号": 1,
        "测试项": "首页短文展示位置",
        "测试步骤": "1. 打开扇贝阅读APP\n2. 进入首页\n3. 观察短文内容的展示位置",
        "期望结果": "短文卡片展示在首页推荐模块第二位",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 2,
        "测试项": "首页短文卡片样式",
        "测试步骤": "1. 打开扇贝阅读APP\n2. 进入首页\n3. 观察短文卡片的展示样式",
        "期望结果": "短文卡片包含标题、类别标签、文章缩略图、难度标签和估计阅读时间",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 3,
        "测试项": "短文内容类别展示",
        "测试步骤": "1. 打开扇贝阅读APP\n2. 进入首页\n3. 观察多个短文卡片\n4. 记录展示的内容类别",
        "期望结果": "短文内容类别符合优化策略，包含新闻、文化、科技等多样化内容",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 4,
        "测试项": "短文难度分布",
        "测试步骤": "1. 打开扇贝阅读APP\n2. 进入首页\n3. 查看多篇短文的难度标签",
        "期望结果": "短文难度分布合理，初级、中级、高级文章都有展示",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 5,
        "测试项": "短文内容更新频率",
        "测试步骤": "1. 记录首页短文内容\n2. 等待系统设定的更新时间\n3. 重新进入APP查看首页短文",
        "期望结果": "短文内容按照系统设定时间更新，旧内容被新内容替换",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 6,
        "测试项": "点击短文卡片跳转",
        "测试步骤": "1. 打开扇贝阅读APP\n2. 进入首页\n3. 点击短文卡片",
        "期望结果": "成功跳转到短文阅读页面，显示完整文章内容",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 7,
        "测试项": "短文阅读页面功能",
        "测试步骤": "1. 打开扇贝阅读APP\n2. 进入首页\n3. 点击短文卡片进入阅读页面\n4. 测试单词查询、笔记等功能",
        "期望结果": "阅读页面的单词查询、笔记、收藏等功能正常可用",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 8,
        "测试项": "用户历史记录影响",
        "测试步骤": "1. 使用新注册账号登录APP\n2. 阅读几篇特定类别的短文\n3. 退出并重新进入APP\n4. 观察首页短文推荐内容",
        "期望结果": "首页短文推荐内容应结合用户阅读历史进行一定程度的个性化推荐",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 9,
        "测试项": "不同设备兼容性",
        "测试步骤": "1. 在不同尺寸的iOS/Android设备上打开APP\n2. 进入首页\n3. 观察短文卡片的展示效果",
        "期望结果": "短文卡片在不同设备上都能正常显示，布局自适应不同屏幕尺寸",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 10,
        "测试项": "网络异常情况",
        "测试步骤": "1. 打开扇贝阅读APP\n2. 切断网络连接\n3. 进入首页\n4. 恢复网络连接后下拉刷新",
        "期望结果": "网络异常时显示之前缓存的内容或友好的错误提示，网络恢复后能正常加载最新内容",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 11,
        "测试项": "短文阅读完成统计",
        "测试步骤": "1. 打开扇贝阅读APP\n2. 进入首页并点击短文\n3. 完整阅读一篇短文\n4. 查看个人阅读统计数据",
        "期望结果": "阅读完成后，个人统计数据更新，包括阅读时长、文章数等指标",
        "测试结果": "",
        "备注": ""
    },
    {
        "序号": 12,
        "测试项": "算法推荐效果A/B测试",
        "测试步骤": "1. 划分用户组：A组使用新算法，B组使用旧算法\n2. 收集两组用户的阅读完成率、停留时间等数据\n3. 对比分析数据",
        "期望结果": "新算法组(A组)在阅读完成率、用户停留时间等关键指标上优于对照组(B组)",
        "测试结果": "",
        "备注": ""
    }
]

# 填充测试用例数据
for row_idx, test_case in enumerate(test_cases, 2):  # 从第2行开始（第1行是标题）
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.value = test_case[header]
        cell.border = border
        if col_idx in [3, 4]:  # 测试步骤和期望结果列
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# 保存Excel文件
file_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "阅读APP首页短文展示策略优化测试用例.xlsx")
wb.save(file_path)
print(f"测试用例已保存至: {file_path}") 